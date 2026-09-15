"""Document loaders for docx/pdf/txt/md/json/csv.

Each loader returns a list of (text, page, section) blocks so page/section
metadata is preserved for citations. Loaders fail safely with IngestionError.
Path security and size checks happen in ingestion before a loader runs.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from .errors import IngestionError, UnsupportedDocument

# (text, page|None, section|None)
Block = tuple[str, int | None, str | None]

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".txt", ".md", ".markdown", ".json", ".csv"}

MIME_BY_SUFFIX = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".json": "application/json",
    ".csv": "text/csv",
}


def load_blocks(path: Path) -> tuple[list[Block], int | None]:
    """Return (blocks, page_count). page_count is None when not applicable."""

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocument(f"unsupported document type: {suffix!r}")

    if suffix == ".docx":
        return _load_docx(path), None
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix in (".txt", ".md", ".markdown"):
        return _load_text(path), None
    if suffix == ".json":
        return _load_json(path), None
    if suffix == ".csv":
        return _load_csv(path), None
    raise UnsupportedDocument(f"unsupported document type: {suffix!r}")  # pragma: no cover


def _load_text(path: Path) -> list[Block]:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise IngestionError(f"file is not valid UTF-8: {path}") from exc
    if not text.strip():
        raise IngestionError(f"empty document: {path}")
    # split markdown by headings into sections when possible
    blocks: list[Block] = []
    current_section = None
    buff: list[str] = []
    for line in text.splitlines():
        if line.startswith("#"):
            if buff:
                blocks.append(("\n".join(buff).strip(), None, current_section))
                buff = []
            current_section = line.lstrip("#").strip()
        buff.append(line)
    if buff:
        blocks.append(("\n".join(buff).strip(), None, current_section))
    return [b for b in blocks if b[0]]


def _load_docx(path: Path) -> list[Block]:
    try:
        import docx
    except Exception as exc:  # pragma: no cover
        raise IngestionError("python-docx not available") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise IngestionError(f"not a valid .docx: {path}: {exc}") from exc
    blocks: list[Block] = []
    section = None
    for para in document.paragraphs:
        if not para.text.strip():
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if "heading" in style:
            section = para.text.strip()
        blocks.append((para.text.strip(), None, section))
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append((" | ".join(cells), None, "table"))
    if not blocks:
        raise IngestionError(f"empty document: {path}")
    return blocks


def _load_pdf(path: Path) -> tuple[list[Block], int]:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover
        raise IngestionError("pypdf not available") from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise IngestionError(f"not a valid PDF: {path}: {exc}") from exc
    blocks: list[Block] = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        text = text.strip()
        if text:
            blocks.append((text, i + 1, None))
    if not blocks:
        raise IngestionError(f"no extractable text in PDF: {path}")
    return blocks, len(reader.pages)


def _load_json(path: Path) -> list[Block]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise IngestionError(f"invalid JSON: {path}: {exc}") from exc
    text = json.dumps(data, indent=2, ensure_ascii=False)
    if not text.strip():
        raise IngestionError(f"empty JSON: {path}")
    return [(text, None, None)]


def _load_csv(path: Path) -> list[Block]:
    blocks: list[Block] = []
    try:
        with open(path, newline="", encoding="utf-8") as fh:
            reader = csv.reader(fh)
            rows = list(reader)
    except (csv.Error, UnicodeDecodeError) as exc:
        raise IngestionError(f"invalid CSV: {path}: {exc}") from exc
    if not rows:
        raise IngestionError(f"empty CSV: {path}")
    header = rows[0]
    for row in rows[1:]:
        pairs = [f"{h}: {v}" for h, v in zip(header, row)]
        if pairs:
            blocks.append(("; ".join(pairs), None, None))
    if not blocks:
        # header-only file: keep the header as one block
        blocks.append((", ".join(header), None, None))
    return blocks
