"""End-to-end pipeline tests: orchestrator, engine, verification, recovery.

These run the REAL pipeline against a REAL .docx fixture. Success is asserted
against actual saved content, not merely the absence of exceptions.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from autoflow_ai.model_gateway import LocalRuleProvider, ModelRegistry, ModelRouter, NoEligibleModel
from autoflow_ai.model_gateway.local_provider import default_local_model
from autoflow_ai.orchestrator import AutoFlow
from autoflow_ai.runtime import DocumentSession, ExecutionEngine, ToolRegistry
from autoflow_ai.runtime.document_tools import register_document_tools
from autoflow_ai.schemas.enums import ExecutionStatus, ModelRole, VerificationStatus
from autoflow_ai.schemas.models import ModelRequest

pytestmark = pytest.mark.integration

FIXTURE = Path(__file__).parent / "fixtures" / "sample.docx"
EM = "\u2014"
PROMPT = (
    "Edit this document by fixing the known wording issue, normalize em-dashes, "
    "remove double spaces, and save the edited version to the same file. "
    'Also replace "teh" with "the".'
)


@pytest.fixture
def work_doc(tmp_path) -> Path:
    dest = tmp_path / "golden.docx"
    shutil.copy(FIXTURE, dest)
    return dest


# ---- GOLDEN END-TO-END ACCEPTANCE TEST -------------------------------------


def test_golden_edit_and_save_same_file(work_doc):
    assert work_doc.exists()  # INPUT FILE EXISTS

    app = AutoFlow.build()
    report = app.run(PROMPT, target_path=str(work_doc))

    # EXECUTION COMPLETE (and only because verification passed)
    assert report.ok, report.error
    assert report.status == ExecutionStatus.COMPLETE

    # PLAN VALID: 4 steps in order
    order_trace = [t for t in report.trace if t.label == "PLAN VALIDATED"]
    assert order_trace and "step_inspect -> step_edit -> step_save -> step_verify" in order_trace[0].detail

    # VERIFICATION PASSED
    assert report.verifications
    assert report.verifications[-1].status == VerificationStatus.PASSED
    assert all(c.passed for c in report.verifications[-1].checks)

    # SAME FILE PATH USED + FILE SAVED
    save_results = [r for r in report.tool_results if r["tool_name"] == "document.save" and r["ok"]]
    assert save_results
    assert save_results[-1]["output"]["same_file"] is True
    assert save_results[-1]["output"]["changed"] is True

    # FILE REOPENED + EXPECTED CONTENT PRESENT + DOCUMENT VALID
    reopened = docx.Document(str(work_doc))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "This document has" in text        # double space removed
    assert "stray word." in text              # space-before-punct fixed
    assert f"plan {EM} and" in text           # em-dash normalized
    assert "review the figures" in text       # teh -> the
    assert "teh" not in text


def test_execution_trace_is_produced(work_doc):
    app = AutoFlow.build()
    report = app.run(PROMPT, target_path=str(work_doc))
    labels = [t.label for t in report.trace]
    for expected in [
        "TASK CREATED",
        "PLAN GENERATED",
        "PLAN VALIDATED",
        "TOOL: document.inspect",
        "TOOL: document.edit",
        "TOOL: document.save",
        "VERIFICATION STARTED",
        "VERIFICATION PASSED",
        "EXECUTION COMPLETE",
    ]:
        assert expected in labels, f"missing trace step: {expected}"


def test_events_can_rebuild_status(work_doc):
    app = AutoFlow.build()
    report = app.run(PROMPT, target_path=str(work_doc))
    event_types = [str(e.type) for e in report.events]
    assert "task.created" in event_types
    assert "plan.validated" in event_types
    assert "verification.passed" in event_types
    assert "execution.completed" in event_types


def test_repeated_execution_does_not_corrupt(work_doc):
    app = AutoFlow.build()
    for _ in range(3):
        report = app.run(PROMPT, target_path=str(work_doc))
        assert report.ok
    text = "\n".join(p.text for p in docx.Document(str(work_doc)).paragraphs)
    assert "Sample Report" in text
    assert "teh" not in text


def test_missing_file_fails_without_complete(tmp_path):
    app = AutoFlow.build()
    report = app.run(PROMPT, target_path=str(tmp_path / "ghost.docx"))
    assert not report.ok
    assert report.status == ExecutionStatus.FAILED
    # never reached COMPLETE
    assert not any(t.label == "EXECUTION COMPLETE" for t in report.trace)


def test_txt_end_to_end(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("hello  world -- ok . teh end", encoding="utf-8")
    app = AutoFlow.build()
    report = app.run(PROMPT, target_path=str(p))
    assert report.ok
    text = p.read_text(encoding="utf-8")
    assert "hello world" in text
    assert f"world {EM} ok." in text
    assert "the end" in text


# ---- VERIFICATION FAILURE PREVENTS COMPLETE --------------------------------


def test_verification_failure_prevents_complete(tmp_path, monkeypatch):
    """Force the reopen/verify to fail and assert we never COMPLETE."""

    dest = tmp_path / "golden.docx"
    shutil.copy(FIXTURE, dest)

    from autoflow_ai.runtime import engine as engine_mod

    real_open = engine_mod.open_document

    def _boom(path):
        # fail only during the verify reopen
        raise engine_mod.DocumentError("forced reopen failure")

    monkeypatch.setattr(engine_mod, "open_document", _boom)

    app = AutoFlow.build()
    report = app.run(PROMPT, target_path=str(dest))
    assert not report.ok
    assert report.status == ExecutionStatus.FAILED
    assert report.verifications[-1].status == VerificationStatus.FAILED


# ---- INVALID PLAN / NO TARGET ----------------------------------------------


def test_no_target_file_is_rejected():
    app = AutoFlow.build()
    from autoflow_ai.brain.planner import PlanError

    with pytest.raises(PlanError):
        app.plan("just fix the grammar please")  # no file anywhere


# ---- MODEL GATEWAY ----------------------------------------------------------


def test_router_selects_local_model():
    reg = ModelRegistry()
    reg.register(default_local_model(), LocalRuleProvider())
    router = ModelRouter(reg)
    req = ModelRequest(
        request_id="mreq_x",
        required_role=ModelRole.PLANNER,
        required_capabilities={"structured_output": True},
    )
    resp = router.generate(req, prompt="hello")
    assert resp.provider == "local-rule"
    assert resp.structured_output["acknowledged"] is True


def test_router_no_eligible_model_raises():
    reg = ModelRegistry()
    reg.register(default_local_model(), LocalRuleProvider())
    router = ModelRouter(reg)
    req = ModelRequest(request_id="mreq_x", required_role=ModelRole.VISION)
    with pytest.raises(NoEligibleModel):
        router.generate(req, prompt="hello")


# ---- BOUNDED RECOVERY -------------------------------------------------------


def test_recovery_is_bounded(tmp_path):
    """A tool that always raises a retryable error must stop after the bound."""

    dest = tmp_path / "work.docx"
    shutil.copy(FIXTURE, dest)

    reg = ToolRegistry()
    session = DocumentSession()
    register_document_tools(reg, session)

    calls = {"n": 0}
    from autoflow_ai.schemas.tools import ToolDefinition
    from autoflow_ai.schemas.enums import RiskClass
    from autoflow_ai.runtime.registry import ToolExecutionError

    def _flaky(args):
        calls["n"] += 1
        raise ToolExecutionError("transient", "temporary glitch")

    reg.register(
        ToolDefinition(
            name="test.flaky",
            description="always fails transiently",
            input_schema={"type": "object", "properties": {}, "required": []},
            risk_class=RiskClass.LOW,
            permission_scope="files:read",
        ),
        _flaky,
    )

    from autoflow_ai.schemas.tasks import TaskGraph, TaskNode
    from autoflow_ai.schemas.enums import AgentKind

    plan = TaskGraph(
        task_id="task_r1",
        goal="recovery bound",
        nodes=(
            TaskNode(
                step_id="step_flaky",
                objective="run flaky tool",
                assigned_agent=AgentKind.DOCUMENT,
                expected_state={"tool": "test.flaky"},
            ),
        ),
    )
    engine = ExecutionEngine(
        registry=reg,
        session=session,
        organization_id="org_local",
        workspace_id="ws_local",
        permissions=frozenset({"files:read"}),
        max_recovery_attempts=2,
    )
    report = engine.run(plan, execution_id="exec_r1")
    assert not report.ok
    # initial attempt + 2 recovery attempts = 3, then stop (not infinite)
    assert calls["n"] == 3


def test_permission_denied_blocks_tool(tmp_path):
    dest = tmp_path / "work.docx"
    shutil.copy(FIXTURE, dest)
    reg = ToolRegistry()
    session = DocumentSession()
    register_document_tools(reg, session)

    from autoflow_ai.schemas.tasks import TaskGraph, TaskNode
    from autoflow_ai.schemas.enums import AgentKind

    plan = TaskGraph(
        task_id="task_p1",
        goal="permission",
        nodes=(
            TaskNode(
                step_id="step_inspect",
                objective="inspect",
                assigned_agent=AgentKind.DOCUMENT,
                expected_state={"tool": "document.inspect", "path": str(dest)},
            ),
        ),
    )
    engine = ExecutionEngine(
        registry=reg,
        session=session,
        organization_id="org_local",
        workspace_id="ws_local",
        permissions=frozenset(),  # no files:read grant
    )
    report = engine.run(plan, execution_id="exec_p1")
    assert not report.ok
    assert "permission denied" in (report.error or "")
