"""Tests for document adapters (docx/txt/md): open, edit, save, reopen, hash."""

from __future__ import annotations

import shutil
import stat
import sys
from pathlib import Path

import docx
import pytest

from autoflow_ai.documents.base import DocumentError, open_document, sha256_file
from autoflow_ai.editing.operations import EditKind, EditOperation

pytestmark = pytest.mark.integration

FIXTURE = Path(__file__).parent / "fixtures" / "sample.docx"
EM = "\u2014"


@pytest.fixture
def docx_copy(tmp_path) -> Path:
    dest = tmp_path / "work.docx"
    shutil.copy(FIXTURE, dest)
    return dest


def _edit_ops():
    return [
        EditOperation(kind=EditKind.REMOVE_DOUBLE_SPACES),
        EditOperation(kind=EditKind.FIX_SPACE_BEFORE_PUNCT),
        EditOperation(kind=EditKind.NORMALIZE_EM_DASHES),
        EditOperation(kind=EditKind.REPLACE_TEXT, find="teh", replace="the"),
    ]


# 1. valid DOCX can be opened
def test_valid_docx_opens(docx_copy):
    adapter = open_document(docx_copy)
    assert adapter is not None


# 2. text can be extracted
def test_docx_text_extracted(docx_copy):
    adapter = open_document(docx_copy)
    text = adapter.read_text()
    assert "Sample Report" in text
    assert "double space" in text


# 3 + 4 + 5 + 6. edit applied, same file saved, reopened, change present
def test_docx_edit_save_reopen_change_present(docx_copy):
    adapter = open_document(docx_copy)
    adapter.apply_operations(_edit_ops())
    saved = adapter.save()  # same file
    assert saved == docx_copy

    reopened = open_document(docx_copy)
    text = reopened.read_text()
    assert "This document has" in text          # double space removed
    assert "stray word." in text                 # space-before-punct fixed
    assert f"plan {EM} and" in text              # em-dash normalized
    assert "review the figures" in text          # teh -> the


# 7. content hash changes when content changes
def test_hash_changes_on_edit(docx_copy):
    before = sha256_file(docx_copy)
    adapter = open_document(docx_copy)
    adapter.apply_operations(_edit_ops())
    adapter.save()
    after = sha256_file(docx_copy)
    assert before != after


# 8. no-op edit does not corrupt document
def test_noop_edit_keeps_document_valid(docx_copy):
    adapter = open_document(docx_copy)
    adapter.apply_operations([EditOperation(kind=EditKind.REPLACE_TEXT, find="zzz-not-present", replace="x")])
    adapter.save()
    reopened = open_document(docx_copy)
    assert "Sample Report" in reopened.read_text()


# 9. malformed DOCX fails safely
def test_malformed_docx_fails_safely(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"this is not a real docx package")
    with pytest.raises(DocumentError):
        open_document(bad)


# 10. read-only / locked file fails safely on save
@pytest.mark.skipif(sys.platform != "win32", reason="read-only semantics vary by OS")
def test_readonly_file_save_fails_safely(docx_copy):
    adapter = open_document(docx_copy)
    adapter.apply_operations(_edit_ops())
    docx_copy.chmod(stat.S_IREAD)
    try:
        with pytest.raises(DocumentError):
            adapter.save()
    finally:
        docx_copy.chmod(stat.S_IWRITE)


def test_missing_file_fails_safely(tmp_path):
    with pytest.raises(DocumentError):
        open_document(tmp_path / "nope.docx")


def test_unsupported_type_fails_safely(tmp_path):
    p = tmp_path / "data.bin"
    p.write_bytes(b"x")
    with pytest.raises(DocumentError):
        open_document(p)


# text/markdown adapters
def test_txt_edit_roundtrip(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("hello  world -- ok .", encoding="utf-8")
    adapter = open_document(p)
    adapter.apply_operations(_edit_ops())
    adapter.save()
    assert p.read_text(encoding="utf-8") == f"hello world {EM} ok."


def test_md_supported(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text("# title\n\ntext  here", encoding="utf-8")
    adapter = open_document(p)
    adapter.apply_operations([EditOperation(kind=EditKind.REMOVE_DOUBLE_SPACES)])
    adapter.save()
    assert "text here" in p.read_text(encoding="utf-8")


# 17. repeated execution does not corrupt the document (idempotent-ish)
def test_repeated_edits_stay_valid(docx_copy):
    for _ in range(3):
        adapter = open_document(docx_copy)
        adapter.apply_operations(_edit_ops())
        adapter.save()
    reopened = open_document(docx_copy)
    text = reopened.read_text()
    assert "Sample Report" in text
    # after the first pass, "teh" is gone and stays gone
    assert "teh" not in text


def test_docx_preserves_paragraph_count(docx_copy):
    before = len(docx.Document(str(docx_copy)).paragraphs)
    adapter = open_document(docx_copy)
    adapter.apply_operations(_edit_ops())
    adapter.save()
    after = len(docx.Document(str(docx_copy)).paragraphs)
    assert before == after  # structure preserved, not rebuilt
