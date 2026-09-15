"""Generate the sample.docx fixture with known, fixable issues.

Run: python tests/fixtures/make_sample.py
Produces tests/fixtures/sample.docx containing:
  * a double space  ("this  has")
  * a space before punctuation ("word .")
  * an ASCII em-dash usage ("A -- B")
  * a known wording token ("teh") for replace tests
Kept deterministic so tests can assert exact post-edit content.
"""

from __future__ import annotations

from pathlib import Path

import docx


def build(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Sample Report", level=1)
    document.add_paragraph("This  document has a double space and a stray word .")
    document.add_paragraph("The plan -- and the result -- should be aligned.")
    document.add_paragraph("Please review teh figures before sending.")
    document.save(str(path))


if __name__ == "__main__":
    out = Path(__file__).parent / "sample.docx"
    build(out)
    print(f"wrote {out}")
